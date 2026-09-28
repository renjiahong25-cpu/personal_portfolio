# -*- coding: utf-8 -*-
r"""把已采用/已生成的规划方案渲染成交付文档（md + docx + pdf）。

用法：
  python scripts\render_plan_doc.py --batch data\output\plan_batch_ai.json --venue-id amap:XXX
  python scripts\render_plan_doc.py --batch data\output\plan_batch_ai.json --venue-name 田涧小苑

输出：data/output/方案_{场地}_{kind}_{时间戳}.md / .docx / .pdf
（pdf 依赖 Word COM，失败不影响 md/docx；最终打印 JSON 结果）
"""
from __future__ import annotations

import argparse
import json
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
OUT = ROOT / "data" / "output"

DAYS_LABEL = {"1": "单日往返", "2": "两天一晚"}
MEETING_LABEL = {"none": "无会议", "half": "半天会议", "full": "全天会议"}
BUDGET_LABEL = {"low": "经济", "mid": "舒适", "high": "轻奢"}
ENT_LABEL = {"rafting": "漂流", "karting": "卡丁车", "boating": "游船", "other": "其他"}
SLOT_ICON = {"drive": "🚗", "meet": "🗣", "meal": "🍽", "activity": "🎯", "hotel": "🛏", "free": "🌴", "return": "🔙"}


def summarize_answers(a: dict) -> str:
    if not a:
        return ""
    bits = []
    if a.get("depart_date"):
        bits.append(f"出发 {a['depart_date']}")
    if a.get("depart_time"):
        bits.append(f"{a['depart_time']} 出发")
    if a.get("days"):
        bits.append(DAYS_LABEL.get(str(a["days"]), str(a["days"])))
    if a.get("people"):
        bits.append(f"{a['people']} 人")
    if a.get("meeting"):
        bits.append(MEETING_LABEL.get(a["meeting"], a["meeting"]))
    ent = [ENT_LABEL.get(x, x) for x in (a.get("ent_prefs") or [])]
    if ent:
        bits.append("娱乐：" + "/".join(ent))
    if a.get("budget"):
        bits.append(BUDGET_LABEL.get(a["budget"], a["budget"]))
    return " · ".join(bits)


def match_activity(slot: dict, detail: dict | None) -> dict | None:
    """按 slot 标签里的 POI 名匹配攻略（精确 → 双向子串模糊）。"""
    if not detail:
        return None
    acts = detail.get("activities") or {}
    label = slot.get("label") or ""
    name = label.split("·")[-1].strip()
    if name in acts:
        return acts[name]
    if name:
        for k, v in acts.items():
            if k and (k in label or name in k or k in name):
                return v
    return None


def detail_slot_md(slot: dict, detail: dict | None) -> list:
    a = match_activity(slot, detail)
    if not a:
        return []
    out: list[str] = []
    if a.get("how_to"):
        out.append(f"> 💡 {a['how_to']}")
    for t in a.get("tips") or []:
        out.append(f"> · {t}")
    ps = a.get("photo_spots") or []
    if ps:
        out.append(f"> 📸 打卡点：{' / '.join(ps)}")
    srcs = a.get("source") or []
    if srcs:
        out.append("> 📕 参考：" + "；".join((s.get("title") or "") for s in srcs[:3]))
    if out:
        out.append("")
    return out


def render_md(name: str, kind: str, answers: dict, plan: dict, detail: dict | None = None) -> str:
    lines = [f"# {name} · 团建行程方案（{kind} 线）", ""]
    s = summarize_answers(answers)
    if s:
        lines += [f"> 参数：{s}", ""]
    if detail and detail.get("overview"):
        lines += ["## 📕 攻略总述（基于小红书研究重组织）", "", detail["overview"], ""]
    for day in plan.get("days") or []:
        lines += [f"## {day.get('label', '')}（{day.get('date', '')}）", "", "| 时间 | 内容 | 驾车 |", "|---|---|---|"]
        for sl in day.get("slots") or []:
            icon = SLOT_ICON.get(sl.get("type"), "▸")
            content = f"{icon} {sl.get('label', '')}"
            drive = ""
            if sl.get("drive_min"):
                drive = f"{sl['drive_min']} 分钟"
                if sl.get("distance_m"):
                    drive += f" / {sl['distance_m'] / 1000:.1f} km"
            lines.append(f"| {sl.get('start', '')}-{sl.get('end', '')} | {content} | {drive} |")
            lines += detail_slot_md(sl, detail)
        lines.append("")
    lines += detail_md(detail)
    warns = plan.get("warnings") or []
    if warns:
        lines += ["## 提示", ""] + [f"- ⚠ {w}" for w in warns] + [""]
    legs = [l for l in (plan.get("map_legs") or []) if l.get("drive_min")]
    if legs:
        lines += ["## 主要车程", ""]
        for l in legs[:4]:
            lines.append(f"- {l.get('from', '')} → {l.get('to', '')}：约 {l['drive_min']} 分钟 / {(l.get('distance_m') or 0) / 1000:.1f} km")
        lines.append("")
    lines.append(f"---\n*venue-scout 自动生成 · {time.strftime('%Y-%m-%d %H:%M:%S')}*")
    return "\n".join(lines)


def detail_md(detail: dict) -> list:
    if not detail:
        return []
    has_extra = (detail.get("food") or detail.get("pitfalls") or detail.get("alternative_nearby")
                 or detail.get("references") or (detail.get("budget_estimate") or {}).get("per_person_range"))
    if not has_extra:
        return []
    lines: list[str] = []
    hl = detail.get("highlights") or []
    if hl:
        lines += ["## 亮点", ""]
        lines += [f"- **{h.get('title', '')}**（{h.get('when', '')}）：{h.get('why', '')}" for h in hl] + [""]
    food = detail.get("food") or []
    if food:
        lines += ["**吃**", ""]
        lines += [f"- {f.get('name', '')}（{f.get('type', '')}）：{f.get('notes', '')} {('｜' + f['source']) if f.get('source') else ''}" for f in food] + [""]
    pits = detail.get("pitfalls") or []
    if pits:
        lines += ["**避坑**", ""]
        lines += [f"- {p.get('pitfall', '')}：{p.get('advice', '')}" for p in pits] + [""]
    bud = detail.get("budget_estimate") or {}
    if bud:
        lines += ["**预算**", ""]
        if bud.get("per_person_range"):
            lines.append(f"- 人均：{bud['per_person_range']}")
        if bud.get("benchmark"):
            lines.append(f"- 基准：{bud['benchmark']}")
        lines += [f"- {b}" for b in bud.get("breakdown") or []] + [""]
    alts = detail.get("alternative_nearby") or []
    if alts:
        lines += ["**附近替代（攻略来自相邻地点，已标注适配）**", ""]
        lines += [f"- {a.get('name', '')}（{a.get('distance', '')}）：{a.get('note', '')}" for a in alts] + [""]
    refs = detail.get("references") or []
    if refs:
        lines += ["## 参考笔记（小红书）", ""]
        for r in refs:
            line = f"- {r.get('title', '')}（{r.get('relevance', '')}）"
            if r.get("url"):
                line += f" {r['url']}"
            lines.append(line)
        lines.append("")
    if detail.get("generated_by"):
        lines += [f"*{detail['generated_by']}*", ""]
    return lines


def detail_slot_docx(doc, slot: dict, detail: dict | None) -> None:
    a = match_activity(slot, detail)
    if not a:
        return
    if a.get("how_to"):
        doc.add_paragraph(f"💡 {a['how_to']}", style="List Bullet 2")
    for t in a.get("tips") or []:
        doc.add_paragraph(f"· {t}", style="List Bullet 2")
    ps = a.get("photo_spots") or []
    if ps:
        doc.add_paragraph(f"📸 打卡点：{' / '.join(ps)}", style="List Bullet 2")
    srcs = a.get("source") or []
    if srcs:
        doc.add_paragraph("📕 参考：" + "；".join((s.get("title") or "") for s in srcs[:3]), style="List Bullet 2")


def detail_docx(doc, detail: dict) -> None:
    if not detail:
        return
    has_extra = (detail.get("food") or detail.get("pitfalls") or detail.get("alternative_nearby")
                 or detail.get("references") or (detail.get("budget_estimate") or {}).get("per_person_range"))
    if not has_extra:
        return
    for h in detail.get("highlights") or []:
        doc.add_paragraph(f"✨ {h.get('title', '')}（{h.get('when', '')}）：{h.get('why', '')}", style="List Bullet")
    for f in detail.get("food") or []:
        doc.add_paragraph(f"{f.get('name', '')}（{f.get('type', '')}）：{f.get('notes', '')}", style="List Bullet")
    for p in detail.get("pitfalls") or []:
        doc.add_paragraph(f"{p.get('pitfall', '')}：{p.get('advice', '')}", style="List Bullet")
    bud = detail.get("budget_estimate") or {}
    if bud:
        doc.add_heading("预算", level=2)
        if bud.get("per_person_range"):
            doc.add_paragraph(f"人均：{bud['per_person_range']}", style="List Bullet")
        for b in bud.get("breakdown") or []:
            doc.add_paragraph(b, style="List Bullet")
    for a in detail.get("alternative_nearby") or []:
        doc.add_paragraph(f"{a.get('name', '')}（{a.get('distance', '')}）：{a.get('note', '')}", style="List Bullet")
    for r in detail.get("references") or []:
        line = f"参考：{r.get('title', '')}（{r.get('relevance', '')}）"
        if r.get("url"):
            line += f" {r['url']}"
        doc.add_paragraph(line, style="List Bullet")


def render_docx(name: str, kind: str, answers: dict, plan: dict, out_docx: Path, detail: dict | None = None) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(f"{name} · 团建行程方案（{kind} 线）", level=0)
    s = summarize_answers(answers)
    if s:
        p = doc.add_paragraph(f"参数：{s}")
        for r in p.runs:
            r.font.size = Pt(10)
    if detail and detail.get("overview"):
        doc.add_heading("📕 攻略总述（基于小红书研究重组织）", level=1)
        doc.add_paragraph(detail["overview"])
    for day in plan.get("days") or []:
        doc.add_heading(f"{day.get('label', '')}（{day.get('date', '')}）", level=1)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "时间", "内容", "驾车"
        for sl in day.get("slots") or []:
            row = table.add_row().cells
            row[0].text = f"{sl.get('start', '')}-{sl.get('end', '')}"
            icon = SLOT_ICON.get(sl.get("type"), "▸")
            row[1].text = f"{icon} {sl.get('label', '')}"
            drive = ""
            if sl.get("drive_min"):
                drive = f"{sl['drive_min']} 分钟"
                if sl.get("distance_m"):
                    drive += f" / {sl['distance_m'] / 1000:.1f} km"
            row[2].text = drive
            detail_slot_docx(doc, sl, detail)
    detail_docx(doc, detail)
    warns = plan.get("warnings") or []
    if warns:
        doc.add_heading("提示", level=1)
        for w in warns:
            doc.add_paragraph(f"⚠ {w}", style="List Bullet")
    legs = [l for l in (plan.get("map_legs") or []) if l.get("drive_min")]
    if legs:
        doc.add_heading("主要车程", level=1)
        for l in legs[:4]:
            doc.add_paragraph(f"{l.get('from', '')} → {l.get('to', '')}：约 {l['drive_min']} 分钟 / {(l.get('distance_m') or 0) / 1000:.1f} km", style="List Bullet")
    doc.save(str(out_docx))


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--batch", required=True)
    ap.add_argument("--venue-id", default="")
    ap.add_argument("--venue-name", default="")
    a = ap.parse_args()

    batch_path = Path(a.batch)
    if not batch_path.exists():
        print(json.dumps({"ok": False, "error": f"未找到 {batch_path}"}, ensure_ascii=False))
        return 1
    d = json.loads(batch_path.read_text(encoding="utf-8"))
    kind = d.get("kind") or "auto"
    answers = d.get("answers") or {}
    item = None
    for it in d.get("items") or []:
        if (a.venue_id and str(it.get("id")) == str(a.venue_id)) or (a.venue_name and it.get("name") == a.venue_name):
            item = it
            break
    if item is None and a.venue_id:
        for it in d.get("items") or []:
            if a.venue_id in str(it.get("id")):
                item = it
                break
    if item is None:
        print(json.dumps({"ok": False, "error": "未找到该场地的方案（可能还没生成完）"}, ensure_ascii=False))
        return 1
    plan = item.get("plan")
    if not plan:
        print(json.dumps({"ok": False, "error": item.get("err") or "该场地方案生成失败"}, ensure_ascii=False))
        return 1
    name = item.get("name") or "场地"
    detail = None
    vid = str(item.get("id") or a.venue_id or "").replace(":", "_")
    detail_f = OUT / f"plan_detail_{vid}.json"
    if detail_f.exists():
        try:
            detail = json.loads(detail_f.read_text(encoding="utf-8"))
        except Exception:
            detail = None
    ts = time.strftime("%Y%m%d_%H%M%S")
    safe = "".join(c for c in name if c.isalnum() or c in " _-（）()")[:24]
    base = OUT / f"方案_{safe}_{kind}_{ts}"
    files = []

    md_path = base.with_suffix(".md")
    md_path.write_text(render_md(name, kind, answers, plan, detail), encoding="utf-8")
    files.append(str(md_path))
    docx_path = base.with_suffix(".docx")
    try:
        render_docx(name, kind, answers, plan, docx_path, detail)
        files.append(str(docx_path))
    except Exception as e:
        print(f"[WARN] docx 失败: {e}", file=sys.stderr)
    pdf_path = base.with_suffix(".pdf")
    if docx_path.exists():
        try:
            from scripts.docx_to_pdf import docx_to_pdf

            docx_to_pdf(str(docx_path), str(pdf_path))
            files.append(str(pdf_path))
        except Exception as e:
            print(f"[WARN] pdf 失败: {e}", file=sys.stderr)
    print(json.dumps({"ok": True, "kind": kind, "venue": name, "files": files}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
