"""Convert venues_summary.md into a formatted .xlsx (bold/filled header, auto column width, hyperlinks).

Pipeline mirrors build_summary_docx.py: parse the same venues_summary.md blocks,
then emit two sheets —「推荐清单」overview table and「场地详情」per-venue details.

Usage:
    python scripts\\build_summary_xlsx.py
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from build_summary_docx import ctrip_url_for, extract_links, parse_blocks
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC = ROOT / "data" / "output" / "venues_summary.md"
OUT = ROOT / "data" / "output" / "团建场地推荐清单.xlsx"

ACCENT = "1F6F4B"
ZEBRA = "EEF5F1"
LINK_COLOR = "0B5CD5"

HEADER_FONT = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=ACCENT)
BODY_FONT = Font(name="Microsoft YaHei", size=10, color="21252B")
ZEBRA_FILL = PatternFill("solid", fgColor=ZEBRA)
LINK_FONT = Font(name="Microsoft YaHei", size=10, color=LINK_COLOR, underline="single")
TITLE_FONT = Font(name="Microsoft YaHei", size=14, bold=True, color=ACCENT)
SECTION_FONT = Font(name="Microsoft YaHei", size=11, bold=True, color="21252B")
NOTE_FONT = Font(name="Microsoft YaHei", size=9, color="6B7280")
THIN = Side(style="thin", color="D9DDE3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def disp_width(text) -> int:
    return sum(2 if ord(ch) > 0x2E7F else 1 for ch in str(text or ""))


def autofit(ws, min_w: int = 8, max_w: int = 60) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            w = max(disp_width(line) for line in str(cell.value).splitlines()) + 2
            if w > widths.get(cell.column, 0):
                widths[cell.column] = w
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_w, min(max_w, w))


def put_link(ws, row: int, col: int, url: str, text: str) -> None:
    c = ws.cell(row=row, column=col, value=text)
    c.hyperlink = url
    c.font = LINK_FONT


def put_text(ws, row: int, col: int, text: str, font: Font = BODY_FONT,
             fill: PatternFill | None = None) -> None:
    c = ws.cell(row=row, column=col, value=text)
    c.font = font
    c.alignment = Alignment(vertical="center", wrap_text=False)
    if fill:
        c.fill = fill


def build_overview(ws, parsed: dict) -> None:
    ws.cell(row=1, column=1, value=parsed["title"] or "团建场地推荐清单").font = TITLE_FONT
    r = 2
    for note in parsed["notes"]:
        note = note.replace("**", "")
        ws.cell(row=r, column=1, value=note).font = NOTE_FONT
        r += 1
    tb = parsed["table"] or {"headers": [], "rows": []}
    headers = list(tb["headers"])
    if not headers:
        return
    headers.append("携程详情")
    head_row = r + 1
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=head_row, column=ci, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.border = BORDER
        c.alignment = Alignment(vertical="center")
    for ri, (row_vals, venue) in enumerate(zip(tb["rows"], parsed["venues"])):
        rr = head_row + 1 + ri
        fill = ZEBRA_FILL if ri % 2 == 1 else None
        row_vals = list(row_vals) + ["" if len(row_vals) < len(headers) - 1 else ""] * max(
            0, len(headers) - 1 - len(row_vals))
        for ci, val in enumerate(row_vals, 1):
            put_text(ws, rr, ci, val, fill=fill)
            ws.cell(row=rr, column=ci).border = BORDER
        url = ctrip_url_for(venue)
        if url:
            put_link(ws, rr, len(headers), url, "打开")
            ws.cell(row=rr, column=len(headers)).border = BORDER
            if fill:
                ws.cell(row=rr, column=len(headers)).fill = fill
    last = head_row + max(1, len(tb["rows"]))
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    ws.auto_filter.ref = f"A{head_row}:{get_column_letter(len(headers))}{last}"


def _detail_lines(ws, r: int, key: str, lines: list[str]) -> int:
    for ln in lines:
        links = extract_links(ln)
        rest = ln
        for url, _text in links:
            rest = rest.replace(f"[{_text}]({url})", "")
        rest = rest.strip(" ·,-")
        if key == "直达":
            if links:
                url, text = links[0]
                put_link(ws, r, 2, url, text)
                if rest:
                    put_text(ws, r, 3, rest)
            r += 1
        elif key == "周边可玩":
            put_text(ws, r, 2, "• " + rest)
            if links:
                url, text = links[0]
                put_link(ws, r, 3, url, text)
            r += 1
        elif key == "实拍参考":
            put_text(ws, r, 2, f"实拍图 {len(links)} 张（图片嵌入见 DOCX/PDF 版）")
            r += 1
            break
        else:
            txt = ln.replace("**", "").strip()
            txt = txt.removeprefix("- ")
            inline = ""
            m = re.search(r"\[([^\]]+)\]\(([^)]+)\)", txt)
            if m:
                inline = m.group(2)
                txt = (txt[: m.start()] + m.group(1) + txt[m.end():]).strip("（）() ")
            prefix = "• " if key == "基本信息" else ""
            put_text(ws, r, 2, prefix + txt)
            if inline:
                put_link(ws, r, 3, inline, "打开")
            r += 1
    return r


def build_details(ws, parsed: dict) -> None:
    r = 1
    for idx, venue in enumerate(parsed["venues"], 1):
        c = ws.cell(row=r, column=1, value=venue["title"])
        c.font = Font(name="Microsoft YaHei", size=12, bold=True, color=ACCENT)
        r += 1
        for key in venue["order"]:
            lines = venue["sections"].get(key) or []
            if not lines:
                continue
            ws.cell(row=r, column=1, value=key).font = SECTION_FONT
            r += 1
            if key in ("直达", "周边可玩", "实拍参考", "基本信息"):
                r = _detail_lines(ws, r, key, lines)
            else:
                for ln in lines:
                    put_text(ws, r, 2, ln.replace("**", ""))
                    r += 1
            r += 1


def main() -> None:
    if not SRC.exists():
        sys.exit(f"!! 缺 {SRC}（先跑 scripts\\build_summary_md.py 生成）")
    parsed = parse_blocks(SRC)
    if not parsed["table"]:
        sys.exit("!! venues_summary.md 中未找到总表，abort")
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "推荐清单"
    build_overview(ws1, parsed)
    autofit(ws1)
    ws2 = wb.create_sheet("场地详情")
    build_details(ws2, parsed)
    ws2.column_dimensions["A"].width = 14
    ws2.column_dimensions["B"].width = 60
    ws2.column_dimensions["C"].width = 40
    wb.save(str(OUT))
    print(f"venues: {len(parsed['venues'])}")
    print(f"written: {OUT}")


if __name__ == "__main__":
    main()
